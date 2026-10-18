"""
Offline AI microservice for the web Speech-to-text app.

Wraps the existing GPT4AllAssistant (ai_assistant.py) so the Node backend can
proxy AI requests to it. The local GPT4All model is loaded once and kept warm.
Only the Python standard library is used for HTTP, so it needs nothing beyond
what the desktop app already requires (gpt4all + the .gguf model in models/).

Run:  python ai_service.py
Environment overrides: AI_SERVICE_HOST (default 127.0.0.1), AI_SERVICE_PORT (8756).
"""

from __future__ import annotations

import json
import os
import re
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

# Resolve config.json and the models/ folder regardless of the launch directory.
os.chdir(Path(__file__).resolve().parent)

from ai_assistant import GPT4AllAssistant
from transcriber import AudioTranscriber, DEFAULT_MODEL
from config import config

HOST = os.environ.get("AI_SERVICE_HOST", "127.0.0.1")
PORT = int(os.environ.get("AI_SERVICE_PORT", "8756"))
REPORT_ROWS_FORMAT_VERSION = "2026-08-20-v2"

assistant = GPT4AllAssistant()


class _PriorityGpuLock:
    """Reusable context manager that serializes all GPU work behind one lock.

    The LLM and Whisper share a single ~8GB GPU and cannot both stay resident,
    so entering the lock swaps the required model onto the GPU (via on_acquire).
    High-priority (interactive LLM) requests must not starve behind the
    relentless background transcription: a waiting high-priority request holds
    the gate, which blocks the next transcription from taking the GPU until the
    LLM has had its turn.
    """

    def __init__(self, gpu, gate, high, on_acquire=None):
        self._gpu = gpu
        self._gate = gate
        self._high = high
        # Called once the GPU lock is held, before the request runs. Used to
        # swap which model is resident so only one occupies the ~8GB GPU.
        self._on_acquire = on_acquire

    def __enter__(self):
        if self._high:
            self._gate.acquire()
            self._gpu.acquire()
            self._gate.release()
        else:
            # Yield to any waiting high-priority request, then take the GPU.
            self._gate.acquire()
            self._gate.release()
            self._gpu.acquire()
        if self._on_acquire is not None:
            try:
                self._on_acquire()
            except Exception:
                self._gpu.release()  # never keep the GPU if the swap failed
                raise
        return self

    def __exit__(self, *_exc):
        self._gpu.release()
        return False


_gpu_lock = threading.Lock()
_hp_gate = threading.Lock()

# When both models are small enough to fit the GPU together (e.g. a 3B LLM
# alongside Whisper), keep both resident so no reloads or transcription pauses
# happen. The GPU compute lock still serializes work so they never run at once.
_KEEP_MODELS_RESIDENT = bool(config.get("ai_keep_models_resident", False))


def _swap_to_llm():
    """Free the GPU for the LLM: unload Whisper, then ensure the LLM is loaded.

    Only one model fits on the ~8GB GPU at a time, so the models are swapped on
    demand. Called while the GPU lock is held, so no other GPU work can race.
    When co-residency is enabled, Whisper is kept loaded instead of unloaded.
    """
    global _transcriber
    if not _KEEP_MODELS_RESIDENT:
        if _transcriber is not None and _transcriber.model is not None:
            try:
                _transcriber.unload_model()
            except Exception as exc:  # pragma: no cover - defensive
                print(f"[AI Service] Whisper unload failed: {exc}")
    if assistant.model is None:
        assistant.load()


def _swap_to_whisper():
    """Free the GPU for Whisper by unloading the LLM.

    Whisper itself is loaded lazily by _get_transcriber() inside the handler.
    Called while the GPU lock is held, so no other GPU work can race.
    When co-residency is enabled, the LLM is kept loaded instead of unloaded.
    """
    if _KEEP_MODELS_RESIDENT:
        return
    if assistant.model is not None:
        try:
            assistant.reset_model()
            import gc
            gc.collect()
        except Exception as exc:  # pragma: no cover - defensive
            print(f"[AI Service] LLM unload failed: {exc}")


# GPT4All text generation is not thread-safe; the high-priority lock also
# serializes every model call. Entering it swaps the LLM onto the GPU.
_model_lock = _PriorityGpuLock(_gpu_lock, _hp_gate, high=True, on_acquire=_swap_to_llm)

# Faster-whisper transcriber (lazy-loaded on first /transcribe request).
# Entering the low-priority lock swaps Whisper onto the GPU (unloads the LLM).
_transcriber = None
_transcriber_lock = _PriorityGpuLock(_gpu_lock, _hp_gate, high=False, on_acquire=_swap_to_whisper)


def _get_transcriber():
    global _transcriber
    if _transcriber is None or _transcriber.model is None:
        model_name = str(
            os.environ.get("WHISPER_MODEL")
            or config.get("transcriber_model_name", DEFAULT_MODEL)
            or DEFAULT_MODEL
        )
        tr = AudioTranscriber()
        ok = tr.load_model(model_name, progress_callback=lambda m: print(f"[Whisper] {m}"), use_gpu=True)
        if not ok:
            raise RuntimeError("Failed to load the Whisper model.")
        _transcriber = tr
    return _transcriber


def _warm_up():
    try:
        assistant.load()
    except Exception as exc:  # pragma: no cover - defensive
        print(f"[AI Service] Model warm-up failed: {exc}")


def _center_docx_cell(cell):
    try:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:
        pass


def _justify_docx_cell(cell, line_spacing=1.15):
    try:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing = line_spacing
    except Exception:
        pass


def _apply_docx_font(doc, font_name="Arial"):
    from docx.shared import Pt
    try:
        normal_style = doc.styles["Normal"]
        if normal_style:
            normal_style.font.name = font_name
    except Exception:
        pass
    for idx, paragraph in enumerate(doc.paragraphs):
        if idx == 0:
            continue
        for run in paragraph.runs:
            run.font.name = font_name
            run.font.size = Pt(11)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = font_name
                        run.font.size = Pt(11)


def _fill_report_docx(template_path, report_data, out_path, movements_justified=True):
    """Fill a .docx report template with report_data and save to out_path.

    Ports the desktop app's _save_report_docx_from_template logic: fills the first
    table's Date/Ref/Subject rows and replaces the Time/Movements/Remarks rows."""
    from docx import Document

    doc = Document(str(template_path))
    date_value = report_data.get("date", "Not available")
    ref_no = report_data.get("ref_no", "Not available")
    subject = report_data.get("subject", "Not available")
    rows = report_data.get("rows", []) or []

    if not doc.tables:
        doc.save(out_path)
        return True

    table = doc.tables[0]
    header_row_index = None
    column_roles = []
    for idx, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        if cells and cells[0].lower().startswith("date"):
            if len(row.cells) > 1:
                row.cells[1].text = str(date_value)
            if len(row.cells) > 2:
                row.cells[2].text = str(date_value)
        if cells and cells[0].lower().startswith("ref"):
            if len(row.cells) > 1:
                row.cells[1].text = str(ref_no)
            if len(row.cells) > 2:
                row.cells[2].text = str(ref_no)
        if cells and cells[0].lower().startswith("subject"):
            if len(row.cells) > 1:
                row.cells[1].text = str(subject)
            if len(row.cells) > 2:
                row.cells[2].text = str(subject)
        normalized_headers = [re.sub(r"\s+", " ", cell.lower()).strip() for cell in cells]
        has_time_column = any(header == "time" or header.startswith("time ") for header in normalized_headers)
        has_movement_column = any("movement" in header for header in normalized_headers)
        if has_time_column and has_movement_column:
            header_row_index = idx
            for header in normalized_headers:
                if header == "date" or header.startswith("date "):
                    column_roles.append("date")
                elif header == "time" or header.startswith("time "):
                    column_roles.append("time")
                elif "movement" in header:
                    column_roles.append("movements")
                elif "remark" in header:
                    column_roles.append("remarks")
                else:
                    column_roles.append("")
            break

    if header_row_index is not None:
        while len(table.rows) > header_row_index + 1:
            table._tbl.remove(table.rows[header_row_index + 1]._tr)

        if not rows:
            rows = [{"time": "Not available", "movements": "Not available", "remarks": "Not available"}]

        for row_data in rows:
            cells = table.add_row().cells
            for column_index, role in enumerate(column_roles[:len(cells)]):
                if role == "date":
                    value = row_data.get("date", "")
                elif role == "time":
                    value = row_data.get("time", "")
                elif role == "movements":
                    value = row_data.get("movements", row_data.get("movement", ""))
                elif role == "remarks":
                    value = row_data.get("remarks", "")
                else:
                    value = ""
                cells[column_index].text = str(value or "")

            for column_index, role in enumerate(column_roles[:len(cells)]):
                if role in {"date", "time"}:
                    _center_docx_cell(cells[column_index])
                elif role == "movements" and movements_justified:
                    _justify_docx_cell(cells[column_index], line_spacing=1.15)

    _apply_docx_font(doc, "Arial")
    doc.save(out_path)
    return True


def _build_ops_report_docx(report_data, out_path):
    """Build the report .docx programmatically (no template required).

    Mirrors the desktop app's _export_processed_report layout: an Ops Reference /
    Name / Entity / Date / DO header block followed by a light-blue "Time |
    Movement" table."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    ops_reference = str(report_data.get("ops_reference", "") or "").strip()
    subject_name = str(report_data.get("subject", "") or "").strip()
    date_value = str(report_data.get("date", "") or "").strip()
    do_value = str(report_data.get("do", "") or "").strip()
    rows = report_data.get("rows", []) or []

    document = Document()

    title = document.add_paragraph()
    title.add_run("Ops Reference: ").bold = True
    title.add_run(ops_reference)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    subject = document.add_paragraph()
    subject.add_run("Name / Entity: ").bold = True
    subject.add_run(subject_name)
    subject.alignment = WD_ALIGN_PARAGRAPH.LEFT

    report_date = document.add_paragraph()
    report_date.add_run("Date: ").bold = True
    report_date.add_run(date_value)
    report_date.alignment = WD_ALIGN_PARAGRAPH.LEFT

    if do_value:
        do_heading = document.add_paragraph()
        do_heading.add_run("DO:").bold = True
        do_heading.add_run(f" {do_value}")
        do_heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False

    # Repeat the header row on new pages.
    tbl = table._element
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = parse_xml(r'<w:tblPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"></w:tblPr>')
        tbl.insert(0, tblPr)
    tblPr.append(parse_xml(r'<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    tblRow = table.rows[0]._element
    trPr = tblRow.trPr
    if trPr is None:
        trPr = parse_xml(r'<w:trPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"></w:trPr>')
        tblRow.insert(0, trPr)
    trPr.append(parse_xml(r'<w:tblHeader xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))

    headers = ["Time", "Movement"]
    header_cells = table.rows[0].cells
    header_cells[0].width = Inches(1.5)
    header_cells[1].width = Inches(4.8)
    for i, header in enumerate(headers):
        header_cells[i].text = header
        shading_elm = parse_xml(r'<w:shd {} w:fill="ADD8E6"/>'.format(nsdecls("w")))
        header_cells[i]._element.get_or_add_tcPr().append(shading_elm)
        for paragraph in header_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if paragraph.runs:
                paragraph.runs[0].bold = True

    for row_data in rows:
        row_cells = table.add_row().cells
        row_cells[0].width = Inches(1.5)
        row_cells[1].width = Inches(4.8)
        row_cells[0].text = str(row_data.get("time", "") or "")
        row_cells[1].text = str(row_data.get("movements", row_data.get("movement", "")) or "")
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    document.save(out_path)
    return True


class Handler(BaseHTTPRequestHandler):
    def _send(self, code, payload):
        body = json.dumps(payload).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _read_json(self):
        length = int(self.headers.get("Content-Length", 0) or 0)
        if length <= 0:
            return {}
        try:
            return json.loads(self.rfile.read(length).decode("utf-8"))
        except Exception:
            return {}

    def log_message(self, *args):  # keep the console quiet
        pass

    def do_GET(self):
        if self.path.rstrip("/") == "/health":
            whisper_ready = _transcriber is not None and _transcriber.model is not None
            llm_ready = assistant.model is not None
            # Models are swapped on demand, so "ready" means the service can
            # serve: a model is resident, or the LLM has no fatal load error.
            ready = llm_ready or whisper_ready or (assistant._load_error is None)
            self._send(200, {
                "ok": True,
                "ready": ready,
                "status": assistant.get_status(),
                "whisperReady": whisper_ready,
            })
            return
        self._send(404, {"error": "Not found"})

    def do_POST(self):
        route = self.path.rstrip("/")
        data = self._read_json()
        try:
            if route == "/ask_general":
                prompt = str(data.get("prompt", "") or "").strip()
                if not prompt:
                    self._send(400, {"error": "prompt is required"})
                    return
                history = data.get("history") if isinstance(data.get("history"), list) else []
                mode = "rewrite" if data.get("mode") == "rewrite" else "chat"
                with _model_lock:
                    answer = assistant.ask_general(prompt, history=history, mode=mode)
                self._send(200, {"answer": answer, "status": assistant.get_status()})
                return

            if route == "/ask_logs":
                prompt = str(data.get("prompt", "") or "").strip()
                logs = data.get("logs") or []
                report_meta = data.get("reportMeta") or None
                if not prompt:
                    self._send(400, {"error": "prompt is required"})
                    return
                with _model_lock:
                    answer = assistant.ask_logs(prompt, logs, report_meta)
                self._send(200, {"answer": answer, "status": assistant.get_status()})
                return

            if route == "/transcribe":
                audio_path = str(data.get("path", "") or "").strip()
                audio_b64 = str(data.get("audioB64", "") or "")
                temp_audio_path = None
                try:
                    if audio_b64:
                        import base64
                        import tempfile

                        filename = os.path.basename(str(data.get("filename", "") or "recording.webm"))
                        suffix = os.path.splitext(filename)[1] or ".webm"
                        temp_audio = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
                        try:
                            temp_audio.write(base64.b64decode(audio_b64))
                        finally:
                            temp_audio.close()
                        temp_audio_path = temp_audio.name
                        audio_path = temp_audio_path
                    if not audio_path or not os.path.exists(audio_path):
                        self._send(400, {"error": "audio file not found"})
                        return
                    with _transcriber_lock:
                        transcriber = _get_transcriber()
                        text, _end, _segments = transcriber.transcribe_file(audio_path, language="en")
                    self._send(200, {"text": (text or "").strip()})
                finally:
                    if temp_audio_path:
                        try:
                            os.unlink(temp_audio_path)
                        except OSError:
                            pass
                return

            if route == "/report_summary":
                entries = data.get("entries") or []
                report_meta = data.get("reportMeta") or {}
                prompt = str(data.get("prompt", "") or "").strip() or (
                    "Summarise the occurrences into a flowing SO Report draft narrative."
                )
                if not entries:
                    self._send(400, {"error": "No entries to summarise."})
                    return
                with _model_lock:
                    answer = assistant.ask_logs(prompt, entries, report_meta)
                self._send(200, {"answer": answer, "status": assistant.get_status()})
                return

            if route == "/report_merged":
                entries = data.get("entries") or []
                report_meta = data.get("reportMeta") or {}
                if not entries:
                    self._send(400, {"error": "No entries to build a report from."})
                    return
                with _model_lock:
                    report = assistant.build_ai_merged_report(entries, report_meta)
                self._send(200, report)
                return

            if route == "/report_rows":
                entries = data.get("entries") or []
                report_meta = data.get("reportMeta") or {}
                if not entries:
                    self._send(400, {"error": "No entries to build a report from."})
                    return
                with _model_lock:
                    report = assistant.build_ai_report_rows(entries, report_meta)
                if isinstance(report, dict):
                    report["formatVersion"] = REPORT_ROWS_FORMAT_VERSION
                self._send(200, report)
                return

            if route == "/report_docx":
                template_path = str(data.get("templatePath", "") or "").strip()
                template_b64 = str(data.get("templateB64", "") or "")
                report_data = data.get("reportData") or {}
                import base64
                import tempfile
                src_path = None
                src_is_temp = False
                if template_b64:
                    # Template bytes were sent because the AI service may run on
                    # a different host than the file; materialize them to a temp.
                    src = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
                    try:
                        src.write(base64.b64decode(template_b64))
                    finally:
                        src.close()
                    src_path = src.name
                    src_is_temp = True
                else:
                    if not template_path or not os.path.exists(template_path):
                        self._send(400, {"error": "Report template not found."})
                        return
                    if str(template_path).lower().rsplit(".", 1)[-1] != "docx":
                        self._send(400, {"error": "Template must be a .docx file."})
                        return
                    src_path = template_path
                tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
                tmp.close()
                try:
                    _fill_report_docx(src_path, report_data, tmp.name, movements_justified=True)
                    with open(tmp.name, "rb") as fh:
                        b64 = base64.b64encode(fh.read()).decode("ascii")
                    self._send(200, {"ok": True, "docx": b64})
                finally:
                    cleanup = [tmp.name] + ([src_path] if src_is_temp else [])
                    for p in cleanup:
                        try:
                            os.unlink(p)
                        except OSError:
                            pass
                return

            if route == "/report_docx_builtin":
                report_data = data.get("reportData") or {}
                import base64
                import tempfile
                tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
                tmp.close()
                try:
                    _build_ops_report_docx(report_data, tmp.name)
                    with open(tmp.name, "rb") as fh:
                        b64 = base64.b64encode(fh.read()).decode("ascii")
                    self._send(200, {"ok": True, "docx": b64})
                finally:
                    try:
                        os.unlink(tmp.name)
                    except OSError:
                        pass
                return

            self._send(404, {"error": "Not found"})
        except Exception as exc:  # pragma: no cover - defensive
            self._send(500, {"error": f"AI service error: {exc}"})


def main():
    print(f"[AI Service] Starting on http://{HOST}:{PORT}")
    # Warm up the LLM in the background so the HTTP server (and Whisper
    # transcription) is available immediately instead of waiting for the 3B load.
    print("[AI Service] Warming up LLM in background (the 3B model can take a while)...")
    threading.Thread(target=_warm_up, daemon=True).start()
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
